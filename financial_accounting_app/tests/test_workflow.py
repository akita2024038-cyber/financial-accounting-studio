import io,json,zipfile
from openpyxl import Workbook, load_workbook
from app.excel import recalculate_cached,validate_xlsx
from app.grading import grade_workbooks,PYSHEETGRADER_AVAILABLE

def books():
    task=Workbook();ws=task.active;ws.title='Journal';ws['A1']='Financial Accounting · Double Entry';ws.merge_cells('A1:C1')
    ws['A3']='Assets';ws['B3']=1000;ws['A4']='Liabilities';ws['B4']=400
    ws['A5']='Equity';ws['B5']=None
    ws['A6']='Assets minus liabilities';ws['B6']=None
    ws['B3'].number_format='#,##0.00 "USD"'
    solution=Workbook();sk=solution.active;sk.title='Journal';sk['A1']='Financial Accounting · Double Entry';sk.merge_cells('A1:C1')
    for c in ['A3','B3','A4','B4','A5','A6']:sk[c]=ws[c].value
    sk['B5']='=B3-B4';sk['B6']='=SUM(B3,-B4)'
    def save(x):b=io.BytesIO();x.save(b);return b.getvalue()
    return save(task),save(solution)

def logged_client(c):
    assert c.post('/api/setup',json={'email':'teacher@example.com','password':'sufficiently-long-pass','setup_token':'test-setup-token'}).status_code==200
    assert c.post('/api/login',json={'email':'teacher@example.com','password':'sufficiently-long-pass'}).status_code==200
    return c

def upload(c):
    t,k=books();r=c.post('/api/assignments',data={'title':'Financial Accounting: Equity','topic':'Accounting equation','instructions':'Calculate equity'},files={'task':('Task.xlsx',t),'key':('Key.xlsx',k)})
    assert r.status_code==200,r.text
    a=r.json();assert len(a['inspection']['inputs'])==2;assert len(a['inspection']['questions'])==2
    inputs=[{**x,'min':x['original']-300,'max':x['original']+300,'step':10} for x in a['inspection']['inputs']]
    questions=a['inspection']['questions'];questions[0]['points']=2.5;questions[1]['points']=1.0
    cfg={'inputs':inputs,'questions':questions,'constraints':[]}
    res=c.put(f'/api/assignments/{a["id"]}/config',json=cfg)
    assert res.status_code==200,res.text
    return a['id'],cfg

def test_workflow_end_to_end(client):
    c=logged_client(client);aid,cfg=upload(c)
    r=c.post(f'/api/assignments/{aid}/generate',json={'count':3});assert r.status_code==200,r.text
    assert r.json()['generated']==3
    tasks_only=c.get(f'/api/assignments/{aid}/file/Tasks_Only.zip')
    keys_only=c.get(f'/api/assignments/{aid}/file/Keys_Only.zip')
    assert tasks_only.status_code==200 and keys_only.status_code==200
    with zipfile.ZipFile(io.BytesIO(tasks_only.content)) as tf,zipfile.ZipFile(io.BytesIO(keys_only.content)) as kf:
        assert len(tf.namelist())==3 and len(kf.namelist())==3
        assert all('KEY' not in x for x in tf.namelist())
    rr=c.get(f'/api/assignments/{aid}/file/All_Variants.zip');assert rr.status_code==200
    with zipfile.ZipFile(io.BytesIO(rr.content)) as z:
        names=z.namelist();assert len(names)==7
        assert 'All_Answers.xlsx' in names
        vals=[]
        for i in range(1,4):
            task=load_workbook(io.BytesIO(z.read(f'Variant_{i:03d}_Task.xlsx')))
            key=load_workbook(io.BytesIO(z.read(f'Variant_{i:03d}_KEY.xlsx')),data_only=True)
            vals.append((task.active['B3'].value,task.active['B4'].value))
            assert task.active['B5'].value is None
            assert key.active['B5'].value==task.active['B3'].value-task.active['B4'].value
            assert key.active['B6'].value==task.active['B3'].value-task.active['B4'].value
            assert task.active['B3'].number_format=='#,##0.00 "USD"'
        assert len(set(vals))==3
    assert c.post(f'/api/assignments/{aid}/publish',json={'published':True}).status_code==200
    assert c.put(f'/api/assignments/{aid}/config',json=cfg).status_code==409
    pub=c.get(f'/api/public/assignments/{aid}');assert pub.status_code==200
    reg=c.post('/api/student/register',json={'assignment_id':aid,'student_id':'BMN17000','group':'BMN4-25','full_name':'Student One'})
    assert reg.status_code==200,reg.text
    token=reg.json()['token'];headers={'X-Student-Token':token}
    assert c.get('/api/student/download').status_code==401
    tdata=c.get('/api/student/download',headers=headers);assert tdata.status_code==200
    student_wb=load_workbook(io.BytesIO(tdata.content))
    student_wb.active['B5']='=B3-B4';student_wb.active['B6']='=SUM(B3,-B4)'
    b=io.BytesIO();student_wb.save(b)
    response=c.post('/api/student/submit',files={'file':('submission.xlsx',b.getvalue())},headers=headers)
    assert response.status_code==200,response.text
    result=response.json();assert result['score']==3.5;assert result['maximum']==3.5
    assert result['counts']['correct']==2
    assert c.get('/api/student/results').status_code==401
    assert c.get('/api/student/results',headers=headers).json()[0]['result']['score']==3.5
    assert c.get(f'/api/student/graded-copy/{result["submission_id"]}',headers=headers).status_code==200
    assert c.get(f'/api/assignments/{aid}/grades-export?format=csv').status_code==200
    assert c.get(f'/api/assignments/{aid}/grades-export?format=xlsx').status_code==200
    assert c.get(f'/api/assignments/{aid}/file/Variant_001_KEY.xlsx').status_code==200
    reg2=c.post('/api/student/register',json={'assignment_id':aid,'student_id':'BMN17001','group':'BMN4-25'})
    assert reg2.status_code==200
    other={'X-Student-Token':reg2.json()['token']}
    assert c.get(f'/api/student/graded-copy/{result["submission_id"]}',headers=other).status_code==404
    bad=student_wb;bad.active['B5']='=B3+B4';bad.active['B6']=None;b=io.BytesIO();bad.save(b)
    graded=c.post('/api/student/submit',headers=headers,files={'file':('bad.xlsx',b.getvalue())})
    assert graded.status_code==200,graded.text
    data=graded.json();assert data['score']<3.5
    assert any(row['status']=='missing' for row in data['rows'])
    assert c.post(f'/api/submissions/{result["submission_id"]}/override',json={'question':'Q1','awarded':1.0,'reason':'Teacher rubric override'}).status_code==200
    assert c.post(f'/api/submissions/{result["submission_id"]}/regrade').status_code==200

def test_zip_batch(client):
    c=logged_client(client);aid,_=upload(c);assert c.post(f'/api/assignments/{aid}/generate',json={'count':1}).status_code==200
    c.post(f'/api/assignments/{aid}/publish',json={'published':True})
    r=c.post('/api/student/register',json={'assignment_id':aid,'student_id':'X902','group':'AC2'})
    b=io.BytesIO();z=zipfile.ZipFile(b,'w');t,_=books();z.writestr('X902__work.xlsx',t);z.close()
    resp=c.post(f'/api/assignments/{aid}/grade-zip',files={'file':('work.zip',b.getvalue())})
    assert resp.status_code==200,resp.text
    assert 'score' in resp.json()['results'][0]

def test_reject_bad_input(client):
    c=logged_client(client)
    assert c.post('/api/assignments',data={'title':'Bad'},files={'task':('evil.xlsx',b'notxlsx'),'key':('key.xlsx',b'notxlsx')}).status_code==422
    aid,cfg=upload(c)
    cfg['inputs'][0]['cell']='A1'
    assert c.put(f'/api/assignments/{aid}/config',json=cfg).status_code==422
    assert c.get(f'/api/assignments/{aid}/file/../../original_key.xlsx').status_code in (404,405,422)

def test_recalculate_cached_equations():
    t,k=books();valid=validate_xlsx(k)
    result=recalculate_cached(k)
    wb=load_workbook(io.BytesIO(result),data_only=True)
    assert wb.active['B5'].value==600
    assert wb.active['B6'].value==600

def test_bundled_thirty_variants_are_all_correct():
    from pathlib import Path
    archive=Path(__file__).resolve().parents[1]/'examples'/'demo_30_variants.zip'
    with zipfile.ZipFile(archive) as z:
        assert len(z.namelist())==61
        combos=set()
        for i in range(1,31):
            t=load_workbook(io.BytesIO(z.read(f'Variant_{i:03d}_Task.xlsx')))
            k=load_workbook(io.BytesIO(z.read(f'Variant_{i:03d}_KEY.xlsx')),data_only=True)
            balance=k['Balance Sheet'];income=k['Income Statement']
            combos.add((t['Balance Sheet']['B4'].value,t['Balance Sheet']['B5'].value,t['Income Statement']['B4'].value,t['Income Statement']['B5'].value))
            assert balance['B6'].value==balance['B4'].value-balance['B5'].value
            assert balance['B7'].value==balance['B4'].value
            assert income['B6'].value==income['B4'].value-income['B5'].value
        assert len(combos)==30

def test_pysheetgrader_symbolic_math_and_security():
    from app.grading import symbolic_equivalent, PYSHEETGRADER_AVAILABLE,csv_safe
    assert PYSHEETGRADER_AVAILABLE
    assert symbolic_equivalent('=B3+B4','=B4+B3')
    assert not symbolic_equivalent('=B3+B4','=B3-B4')
    assert not symbolic_equivalent('=__import__("os")','=B3')
    assert csv_safe('=1+2')=="'=1+2"

def test_partial_credit_for_hardcoded_formula(client):
    c=logged_client(client);aid,_=upload(c)
    assert c.post(f'/api/assignments/{aid}/generate',json={'count':1}).status_code==200
    c.post(f'/api/assignments/{aid}/publish',json={'published':True})
    token=c.post('/api/student/register',json={'assignment_id':aid,'student_id':'A77','group':'Finance01'}).json()['token']
    headers={'X-Student-Token':token}
    data=c.get('/api/student/download',headers=headers).content
    wb=load_workbook(io.BytesIO(data));sheet=wb.active
    sheet['B5']=sheet['B3'].value-sheet['B4'].value # hardcoded but correct
    sheet['B6']='=SUM(B3,-B4)'
    out=io.BytesIO();wb.save(out)
    res=c.post('/api/student/submit',headers=headers,files={'file':('partial.xlsx',out.getvalue())})
    assert res.status_code==200,res.text
    rows=res.json()['rows']
    assert rows[0]['status']=='partial'
    assert rows[0]['awarded']==0.625
    assert rows[1]['status']=='correct'
